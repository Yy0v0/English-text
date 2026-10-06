# -*- coding: utf-8 -*-
"""
同步题库.py
──────────
将「题库.xlsx」中「题库」Sheet 的内容同步回「题库数据.js」。

用途：
  在 Excel 中编辑题目后，运行本脚本，js 文件会自动更新，
  然后刷新游戏页面即可看到新题目。

使用方法：
  在命令行运行：python 同步题库.py
  或双击运行本脚本。

依赖：pip install openpyxl
"""

import os
import re
from openpyxl import load_workbook

# ═══════════════════════════════════════════════════════════
#  配置
# ═══════════════════════════════════════════════════════════
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
JS_FILE = os.path.join(SCRIPT_DIR, '题库数据.js')
XLSX_FILE = os.path.join(SCRIPT_DIR, '题库.xlsx')

# ═══════════════════════════════════════════════════════════
#  1. 读取 xlsx
# ═══════════════════════════════════════════════════════════
print('读取 Excel...')
wb = load_workbook(XLSX_FILE, data_only=True)

# 找「题库」Sheet
if '题库' in wb.sheetnames:
    ws = wb['题库']
elif '题库' in [s.title for s in wb.worksheets]:
    ws = wb['题库']
else:
    ws = wb.active
    print(f'WARNING: 未找到「题库」Sheet，使用当前 Sheet: {ws.title}')

# 读取表头（第一行）
headers = [cell.value for cell in ws[1]]
print(f'表头：{headers}')

# 找各列索引
def col(name):
    try:
        return headers.index(name)
    except ValueError:
        return -1

idx_type    = col('type')
idx_diff    = col('diff')
idx_q       = col('q')
idx_options = col('options')
idx_ans     = col('ans')
idx_explain = col('explain')

missing = [n for n, i in [('type',idx_type),('diff',idx_diff),('q',idx_q),
                           ('options',idx_options),('ans',idx_ans),('explain',idx_explain)]
           if i == -1]
if missing:
    print(f'ERROR: 缺少列 {missing}，请检查 Excel 表头')
    exit(1)

# 解析每行
questions = []
for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
    type_v    = str(row[idx_type]).strip()    if row[idx_type]    else ''
    diff_v    = str(row[idx_diff]).strip()    if row[idx_diff]    else ''
    q_v       = str(row[idx_q]).strip()       if row[idx_q]       else ''
    options_v = str(row[idx_options]).strip() if row[idx_options] else ''
    ans_v     = int(row[idx_ans])             if row[idx_ans]     is not None else 0
    explain_v = str(row[idx_explain]).strip() if row[idx_explain] else ''

    if not type_v or not q_v:
        continue  # 跳过空行

    # 解析 options（英文逗号分隔）
    opts = [o.strip() for o in options_v.split(',')]
    opts = [o for o in opts if o]  # 去除空选项
    if len(opts) < 4:
        print(f'WARNING: 第{row_idx}行选项不足4个，跳过')
        continue

    questions.append({
        'type':    type_v,
        'diff':    diff_v,
        'q':       q_v,
        'options': opts,
        'ans':     ans_v,
        'explain': explain_v,
    })

print(f'读取到 {len(questions)} 道题目')

# ═══════════════════════════════════════════════════════════
#  2. 生成 JS 代码片段
# ═══════════════════════════════════════════════════════════
diff_labels = {
    'low':  '小学低年级 (low)',
    'high': '小学高年级 (high)',
    'chu1': '初一 (chu1)',
    'chu2': '初二 (chu2)',
}

def escape_js_str(s):
    """对字符串进行 JS 转义（单引号用\'，反斜杠用\\）"""
    return s.replace('\\', '\\\\').replace("'", "\\'")

lines = ['const QUESTION_BANK = [']
current_diff = None
for q in questions:
    diff = q['diff']
    # 按难度分组加注释
    if diff != current_diff:
        label = diff_labels.get(diff, diff)
        lines.append(f'')
        lines.append(f'  // ========== {label} ==========')
        current_diff = diff

    opts_str = ','.join(f"'{escape_js_str(o)}'" for o in q['options'])
    explain_str = escape_js_str(q['explain'])
    q_str = escape_js_str(q['q'])
    lines.append(
        f"  {{type:'{q['type']}', diff:'{diff}', q:'{q_str}',"
        f"options:[{opts_str}],ans:{q['ans']},explain:'{explain_str}'}},"
    )

lines.append('];')
js_array_text = '\n'.join(lines)

# ═══════════════════════════════════════════════════════════
#  3. 替换 js 文件中的 QUESTION_BANK 数组
# ═══════════════════════════════════════════════════════════
with open(JS_FILE, 'r', encoding='utf-8') as f:
    js_content = f.read()

# 替换 QUESTION_BANK 数组
new_js, count = re.subn(
    r'(const QUESTION_BANK = )\[[\s\S]*?\]\;',
    js_array_text,
    js_content
)

if count == 0:
    print('ERROR: 未找到 const QUESTION_BANK = [...] 进行替换')
    exit(1)

with open(JS_FILE, 'w', encoding='utf-8') as f:
    f.write(new_js)

print(f'已同步到：{JS_FILE}')
print(f'共写入 {len(questions)} 道题目')

# 统计
from collections import Counter
by_diff = Counter(q['diff'] for q in questions)
by_type = Counter(q['type'] for q in questions)
print('按难度：', dict(by_diff))
print('按题型：', dict(by_type))
print('')
print('同步完成！请刷新游戏页面使更改生效。')
