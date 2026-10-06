# -*- coding: utf-8 -*-
"""
build_xlsx.py
────────────
将「题库数据.js」中的题库数据导出为 Excel 表格（题库.xlsx）。

用途：查看题库内容；用 Excel 编辑后运行「同步题库.py」同步回 js。
使用方法：命令行运行 python build_xlsx.py 或双击运行。
依赖：pip install openpyxl
"""

import re, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# ── 辅助函数 ──────────────────────────────────────────────
def parse_question(text):
    def get_str(key):
        m = re.search(rf"{key}:\s*['\"]([^'\"]*)['\"]", text)
        return m.group(1).strip() if m else ''
    def get_arr(key):
        m = re.search(rf"{key}:\s*\[([^\]]*)\]", text)
        if not m: return []
        return [p.strip() for p in re.findall(r"['\"]([^'\"]*)['\"]", m.group(1))]
    def get_int(key):
        m = re.search(rf"{key}:\s*(\d+)", text)
        return int(m.group(1)) if m else 0
    return {
        'type':    get_str('type'),
        'diff':    get_str('diff'),
        'q':       get_str('q'),
        'options': get_arr('options'),
        'ans':     get_int('ans'),
        'explain': get_str('explain'),
    }

# ── 1. 读取题库数据.js ─────────────────────────────────────
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
with open(os.path.join(SCRIPT_DIR, '题库数据.js'), 'r', encoding='utf-8') as f:
    raw = f.read()

m = re.search(r'const QUESTION_BANK = (\[[\s\S]*?\]);', raw)
if not m:
    print('ERROR: 未找到 const QUESTION_BANK = [...]'); exit(1)

array_text = m.group(1)
questions = []
depth, obj_start, i = 0, None, 0
while i < len(array_text):
    ch = array_text[i]
    if ch == '{':
        if depth == 0: obj_start = i
        depth += 1
    elif ch == '}':
        depth -= 1
        if depth == 0 and obj_start is not None:
            q = parse_question(array_text[obj_start:i+1])
            if q: questions.append(q)
            obj_start = None
    i += 1

print(f'读取到 {len(questions)} 道题目')

# ── 2. 构建 Excel ───────────────────────────────────────────
wb = Workbook()

# Sheet1: 说明
ws_n = wb.active; ws_n.title = '说明'
for i, (text, bold, size, color) in enumerate([
    ('题库修改说明',         True,  16, None),
    ('',                     False, 11, None),
    ('★ 双向同步工作流',      True,  12, 'C00000'),
    ('  本表格由「build_xlsx.py」自动从「题库数据.js」导出。', False, 11, None),
    ('  编辑 xlsx 后运行「同步题库.py」→ 将修改同步回 js。', False, 11, None),
    ('',                     False, 11, None),
    ('★ Sheet「题库」各列说明', True, 12, None),
    ('  type    ：题型，填 单词 / 语法 / 翻译', False, 11, None),
    ('  diff    ：难度，填 low / high / chu1 / chu2', False, 11, None),
    ('  q       ：题目文字', False, 11, None),
    ('  options ：四个选项，英文逗号分隔', False, 11, None),
    ('  ans     ：正确答案序号，0=A, 1=B, 2=C, 3=D', False, 11, None),
    ('  explain ：答对后显示的解析文字', False, 11, None),
    ('',                     False, 11, None),
    ('★ 难度对应关系',        True, 12, None),
    ('  low  = 小学低年级（1-3年级）', False, 11, None),
    ('  high = 小学高年级（4-6年级）', False, 11, None),
    ('  chu1 = 初一（七年级）', False, 11, None),
    ('  chu2 = 初二（八年级）', False, 11, None),
], 2):
    c = ws_n.cell(row=i, column=1, value=text)
    c.font = Font(bold=bold, size=size or 11, color=color or '000000')
ws_n.column_dimensions['A'].width = 58

# Sheet2: 题库
ws = wb.create_sheet('题库')
headers = ['type', 'diff', 'q', 'options', 'ans', 'explain']
hdr_colors = {'type':'9B59B6','diff':'3498DB','q':'27AE60','options':'34495E','ans':'FFD700','explain':'95A5A6'}
ws.append(headers)
for col, name in enumerate(headers, 1):
    c = ws.cell(1, col)
    c.fill = PatternFill('solid', start_color=hdr_colors[name])
    c.font = Font(bold=True, color='FFFFFF', size=11)
    c.alignment = Alignment(horizontal='center', vertical='center')
    c.border = Border(left=Side('thin','FFCCCCCC'), right=Side('thin','FFCCCCCC'),
                      top=Side('thin','FFCCCCCC'), bottom=Side('thin','FFCCCCCC'))

for q in questions:
    ws.append([q['type'], q['diff'], q['q'], ','.join(q['options']), q['ans'], q['explain']])

diff_bg = {'low':'FFEBF3FB','high':'FFEDFAF0','chu1':'FFFFF8E1','chu2':'FFFCE4EC'}
bd = Border(left=Side('thin','FFCCCCCC'), right=Side('thin','FFCCCCCC'),
           top=Side('thin','FFCCCCCC'), bottom=Side('thin','FFCCCCCC'))
for row in ws.iter_rows(min_row=2):
    bg = diff_bg.get(row[1].value, 'FFFFFFFF')
    fill = PatternFill('solid', start_color=bg)
    for cell in row:
        cell.fill = fill; cell.border = bd
        if cell.column in (3,4,6): cell.alignment = Alignment(wrap_text=True, vertical='top')

ws.column_dimensions['A'].width = 8
ws.column_dimensions['B'].width = 8
ws.column_dimensions['C'].width = 38
ws.column_dimensions['D'].width = 55
ws.column_dimensions['E'].width = 6
ws.column_dimensions['F'].width = 36
ws.freeze_panes = 'A2'

# ── 3. 保存 ───────────────────────────────────────────────
out = os.path.join(SCRIPT_DIR, '题库.xlsx')
wb.save(out)
print(f'已生成：{out}')
from collections import Counter
print('按难度：', dict(Counter(q['diff'] for q in questions)))
print('按题型：', dict(Counter(q['type'] for q in questions)))
